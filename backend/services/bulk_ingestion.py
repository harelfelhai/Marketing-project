"""
services/bulk_ingestion.py — Phase E1 batch phone insertion.

Simplified to remove scoring service dependency. Uses new PhoneNumber schema.
"""

import csv
import io
import re
import uuid
from typing import Optional

import openpyxl
from openpyxl import Workbook
from sqlalchemy.exc import IntegrityError

from exceptions import TargetNotFoundError
from models.entity import Entity
from models.phone_number import PhoneNumber
from models.types import SOFT_DELETE_SENTINEL, not_deleted


# ---------------------------------------------------------------------------
# Bulk-upload constants
# ---------------------------------------------------------------------------

BULK_UPLOAD_REQUIRED_COLUMNS = (
    "phone_number",
    "entity_id",
    "ingestion_source",
)
BULK_UPLOAD_OPTIONAL_COLUMNS = (
    "phone_type",
    "target_entity_id",
)
BULK_UPLOAD_ALL_COLUMNS = BULK_UPLOAD_REQUIRED_COLUMNS + BULK_UPLOAD_OPTIONAL_COLUMNS

BULK_UPLOAD_MAX_ROWS = 5_000

_TOKEN_SPLIT_RE = re.compile(r"[,\s;]+")
_PHONE_CLEAN_RE = re.compile(r"[^\d+]")
_PHONE_REGEX = re.compile(r"^\+?\d+$")
_FAILURE_INPUT_CAP = 200


def _tokenize(raw: str) -> list[tuple[int, str]]:
    pairs: list[tuple[int, str]] = []
    idx = 0
    for token in _TOKEN_SPLIT_RE.split(raw):
        if not token:
            continue
        idx += 1
        pairs.append((idx, token))
    return pairs


def _normalize(token: str) -> str:
    return _PHONE_CLEAN_RE.sub("", token.strip())


class BulkIngestionService:
    """Sole writer to the bulk-ingestion path."""

    def __init__(self, storage) -> None:
        self.entities = storage.entities
        self.phones = storage.phones

    def ingest_bulk_text(
        self,
        phone_numbers_raw: str,
        entity_id: str,
        ingestion_source: str,
        phone_type: Optional[str] = None,
        extra_shared: Optional[dict] = None,
        phone_extra_shared: Optional[dict] = None,  # deprecated alias
        uploaded_by_user_id: Optional[str] = None,
    ) -> dict:
        """
        Execute two-pass bulk-text ingestion for phones attached to an existing entity.

        Raises:
            TargetNotFoundError: entity_id is invalid or soft-deleted.
        """
        submission_id = str(uuid.uuid4())

        ent = self.entities.get(entity_id)
        if ent is None or ent.deleted_at != SOFT_DELETE_SENTINEL:
            raise TargetNotFoundError(target_phone_number=f"entity_id={entity_id}")

        # Pass 1 — tokenize, normalize, validate, deduplicate.
        failed_rows: list[dict] = []
        candidates: list[tuple[int, str, str]] = []
        seen: dict[str, int] = {}

        for row, raw in _tokenize(phone_numbers_raw):
            normalized = _normalize(raw)
            if not normalized:
                failed_rows.append({"row": row, "input": raw[:_FAILURE_INPUT_CAP], "error": "Empty after normalization"})
                continue
            if not _PHONE_REGEX.match(normalized):
                failed_rows.append({"row": row, "input": raw[:_FAILURE_INPUT_CAP], "error": "Invalid phone format"})
                continue
            if normalized in seen:
                failed_rows.append({"row": row, "input": raw[:_FAILURE_INPUT_CAP], "error": f"Duplicate of row {seen[normalized]} in this batch"})
                continue
            seen[normalized] = row
            candidates.append((row, raw, normalized))

        if not candidates:
            return {
                "success_count": 0, "failed_count": len(failed_rows),
                "phone_ids": [], "entity_ids": [],
                "failed_rows": failed_rows, "bulk_submission_id": submission_id,
            }

        phone_extra_with_audit = dict(extra_shared or phone_extra_shared or {})
        phone_extra_with_audit["bulk_submission_id"] = submission_id

        phone_ids: list[str] = []
        for row, raw, normalized in candidates:
            if self.phones.list({"phone_number": normalized}, limit=1):
                failed_rows.append({"row": row, "input": raw[:_FAILURE_INPUT_CAP], "error": "Already exists in the system"})
                continue
            try:
                phone = PhoneNumber(
                    entity_id=ent.id,
                    phone_number=normalized,
                    ingestion_source=ingestion_source,
                    phone_type=phone_type,
                    score=0.0,
                    deleted_at=not_deleted(),
                    extra_data=dict(phone_extra_with_audit),
                )
                self.phones.add(phone)
                phone_ids.append(phone.id)
            except IntegrityError as exc:
                detail = str(exc.orig) if exc.orig else "Database constraint violation"
                failed_rows.append({
                    "row": row, "input": raw[:_FAILURE_INPUT_CAP],
                    "error": (
                        "Already exists in the system"
                        if "UNIQUE" in detail.upper()
                        else f"DB error: {detail[:120]}"
                    ),
                })

        failed_rows.sort(key=lambda r: r["row"])
        return {
            "success_count": len(phone_ids), "failed_count": len(failed_rows),
            "phone_ids": phone_ids, "entity_ids": [ent.id],
            "failed_rows": failed_rows, "bulk_submission_id": submission_id,
        }

    def ingest_bulk_upload(
        self,
        file_bytes: bytes,
        filename: str,
        uploaded_by_user_id: Optional[str] = None,
    ) -> dict:
        """Parse an Excel/CSV file and insert one PhoneNumber per row."""
        submission_id = str(uuid.uuid4())
        lower = filename.lower()
        if lower.endswith(".csv"):
            rows = self._parse_csv(file_bytes)
        elif lower.endswith(".xlsx"):
            rows = self._parse_xlsx(file_bytes)
        else:
            raise ValueError(f"Unsupported file extension: '{filename}'. Allowed: .xlsx, .csv")

        if len(rows) > BULK_UPLOAD_MAX_ROWS:
            raise ValueError(f"Too many rows: {len(rows)} (max {BULK_UPLOAD_MAX_ROWS})")

        failed_rows: list[dict] = []
        candidates: list[tuple[int, dict]] = []
        seen_phones: dict[str, int] = {}

        for row_idx, row in enumerate(rows, start=1):
            raw_phone = str(row.get("phone_number") or "").strip()
            if not raw_phone:
                failed_rows.append({"row": row_idx, "input": _safe_input_str(row), "error": "Missing phone_number"})
                continue
            normalized = _normalize(raw_phone)
            if not normalized or not _PHONE_REGEX.match(normalized):
                failed_rows.append({"row": row_idx, "input": raw_phone[:_FAILURE_INPUT_CAP], "error": "Invalid phone format"})
                continue
            missing = [c for c in BULK_UPLOAD_REQUIRED_COLUMNS[1:] if not row.get(c) and row.get(c) != 0]
            if missing:
                failed_rows.append({"row": row_idx, "input": raw_phone[:_FAILURE_INPUT_CAP], "error": f"Missing required field(s): {', '.join(missing)}"})
                continue
            if normalized in seen_phones:
                failed_rows.append({"row": row_idx, "input": raw_phone[:_FAILURE_INPUT_CAP], "error": f"Duplicate of row {seen_phones[normalized]} in this batch"})
                continue
            seen_phones[normalized] = row_idx
            row["_normalized_phone"] = normalized
            candidates.append((row_idx, row))

        phone_ids: list[str] = []
        entity_ids: list[str] = []

        for row_idx, row in candidates:
            try:
                entity_id = str(row.get("entity_id") or "").strip()
                if not entity_id:
                    raise ValueError("entity_id is required")
                ent = self.entities.get(entity_id)
                if ent is None or ent.deleted_at != SOFT_DELETE_SENTINEL:
                    raise ValueError(f"entity_id={entity_id} not found or deleted")

                if self.phones.list({"phone_number": row["_normalized_phone"]}, limit=1):
                    raise IntegrityError(statement=None, params=None, orig=Exception("UNIQUE phone_number"))

                phone = PhoneNumber(
                    entity_id=ent.id,
                    phone_number=row["_normalized_phone"],
                    ingestion_source=row["ingestion_source"],
                    phone_type=row.get("phone_type"),
                    score=0.0,
                    deleted_at=not_deleted(),
                    extra_data={"bulk_submission_id": submission_id},
                )
                self.phones.add(phone)
                phone_ids.append(phone.id)
                entity_ids.append(ent.id)
            except (IntegrityError, ValueError) as exc:
                if isinstance(exc, IntegrityError):
                    detail = str(exc.orig) if exc.orig else "Database constraint violation"
                    msg = ("Already exists in the system" if "UNIQUE" in detail.upper() else f"DB error: {detail[:120]}")
                else:
                    msg = str(exc)
                failed_rows.append({"row": row_idx, "input": row.get("_normalized_phone", "")[:_FAILURE_INPUT_CAP], "error": msg})

        failed_rows.sort(key=lambda r: r["row"])
        return {
            "success_count": len(phone_ids), "failed_count": len(failed_rows),
            "phone_ids": phone_ids, "entity_ids": entity_ids,
            "failed_rows": failed_rows, "bulk_submission_id": submission_id,
        }

    @staticmethod
    def _parse_csv(file_bytes: bytes) -> list[dict]:
        try:
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError(f"CSV not UTF-8 decodable: {exc}") from exc
        reader = csv.DictReader(io.StringIO(text))
        if reader.fieldnames is None:
            raise ValueError("CSV is empty or has no header row")
        missing = [c for c in BULK_UPLOAD_REQUIRED_COLUMNS if c not in reader.fieldnames]
        if missing:
            raise ValueError(f"CSV header missing required columns: {', '.join(missing)}")
        return list(reader)

    @staticmethod
    def _parse_xlsx(file_bytes: bytes) -> list[dict]:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError(f"Not a valid .xlsx workbook: {exc}") from exc
        ws = wb.active
        if ws is None:
            raise ValueError("Workbook has no active sheet")
        rows_iter = ws.iter_rows(values_only=True)
        header_row = None
        for row in rows_iter:
            if any(cell is not None and str(cell).strip() for cell in row):
                header_row = row
                break
        if header_row is None:
            raise ValueError("Workbook is empty")
        header = [str(c).strip() if c is not None else "" for c in header_row]
        missing = [c for c in BULK_UPLOAD_REQUIRED_COLUMNS if c not in header]
        if missing:
            raise ValueError(f"Workbook header missing required columns: {', '.join(missing)}")
        result = []
        for row in rows_iter:
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            row_dict: dict = {}
            for col_name, cell in zip(header, row):
                if not col_name:
                    continue
                if isinstance(cell, str):
                    cell = cell.strip() or None
                row_dict[col_name] = cell
            result.append(row_dict)
        return result

    @staticmethod
    def generate_template_xlsx() -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "data"
        ws.append(list(BULK_UPLOAD_ALL_COLUMNS))
        ws.append(["+14155551111", "entity_id_here", "manual", "mobile", None])
        ws.append(["+14155551112", "entity_id_here", "manual", "home",   None])
        ins = wb.create_sheet(title="instructions")
        ins.append(["field", "description", "required"])
        ins.append(["phone_number",    "Phone in E.164 format", "yes"])
        ins.append(["entity_id",       "Entity PK to attach phone to", "yes"])
        ins.append(["ingestion_source","Origin (manual/automated)", "yes"])
        ins.append(["phone_type",      "mobile/home/work/other", "no"])
        ins.append(["target_entity_id","Root entity FK (informational)", "no"])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def _safe_input_str(row: dict) -> str:
    parts = [f"{col}={row.get(col)}" for col in BULK_UPLOAD_ALL_COLUMNS if row.get(col) not in (None, "")]
    rendered = " | ".join(parts)
    return rendered[:_FAILURE_INPUT_CAP] if rendered else "(empty row)"
