"""Service-layer tests for ExportService."""

import io

import openpyxl
import pytest

from models.entity import Entity
from models.phone_number import PhoneNumber
from models.types import not_deleted
from services.export import (
    ALLOWED_EXPORT_COLUMNS_PHONES,
    ALLOWED_EXPORT_COLUMNS_TASKS,
    MAX_EXPORT_ROWS,
    ExportService,
)
from services.tasks import PipelineTaskService
from repositories.storage import SqlStorage


def _col(key, label=None, fmt="text"):
    return {"key": key, "label": label or key, "format": fmt}


def _read_workbook(xlsx_bytes):
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True)


def _data_rows(wb):
    ws = wb["data"]
    return [tuple(c.value for c in row) for row in ws.iter_rows()]


@pytest.fixture()
def svc(session):
    return ExportService(storage=SqlStorage(session))


@pytest.fixture()
def seeded_phones(session):
    """Two phones on one entity, one phone on another."""
    e1 = Entity(relation_type="primary", full_name="Alice", deleted_at=not_deleted())
    e2 = Entity(relation_type="primary", full_name="Bob", deleted_at=not_deleted())
    session.add_all([e1, e2])
    session.flush()
    p1 = PhoneNumber(
        entity_id=e1.id, phone_number="+15550001111",
        ingestion_source="manual", verification_status="pending",
        score=0.0, deleted_at=not_deleted(),
    )
    p2 = PhoneNumber(
        entity_id=e1.id, phone_number="+15550002222",
        ingestion_source="automated", verification_status="verified",
        score=0.8, deleted_at=not_deleted(),
    )
    p3 = PhoneNumber(
        entity_id=e2.id, phone_number="+15550003333",
        ingestion_source="manual", verification_status="pending",
        score=0.0, deleted_at=not_deleted(),
    )
    session.add_all([p1, p2, p3])
    session.commit()
    return [p1, p2, p3]


# ===========================================================================
# Phones export — happy path
# ===========================================================================


class TestExportPhonesHappyPath:
    def test_returns_bytes_and_filename(self, svc, seeded_phones):
        xlsx_bytes, filename = svc.export_phones(
            filters={},
            columns=[_col("phone_number"), _col("verification_status")],
        )
        assert isinstance(xlsx_bytes, bytes)
        assert xlsx_bytes[:2] == b"PK"
        assert filename.startswith("phones_")
        assert filename.endswith(".xlsx")

    def test_header_row_uses_caller_supplied_labels(self, svc, seeded_phones):
        xlsx_bytes, _ = svc.export_phones(
            filters={},
            columns=[
                _col("phone_number", label="Phone Number"),
                _col("verification_status", label="Status"),
            ],
        )
        wb = _read_workbook(xlsx_bytes)
        rows = _data_rows(wb)
        assert rows[0] == ("Phone Number", "Status")

    def test_data_rows_resolve_each_column(self, svc, seeded_phones):
        xlsx_bytes, _ = svc.export_phones(
            filters={},
            columns=[_col("phone_number"), _col("verification_status")],
        )
        wb = _read_workbook(xlsx_bytes)
        rows = _data_rows(wb)
        assert len(rows) == 4  # 1 header + 3 phones
        phone_numbers = {r[0] for r in rows[1:]}
        assert phone_numbers == {"+15550001111", "+15550002222", "+15550003333"}

    def test_every_default_allowlist_key_is_accepted(self, svc, seeded_phones):
        cols = [_col(k) for k in ALLOWED_EXPORT_COLUMNS_PHONES]
        xlsx_bytes, _ = svc.export_phones(filters={}, columns=cols)
        assert xlsx_bytes[:2] == b"PK"


# ===========================================================================
# Filters
# ===========================================================================


class TestExportPhonesFilters:
    def test_verification_status_filter_narrows_rows(self, svc, seeded_phones):
        xlsx_bytes, _ = svc.export_phones(
            filters={"verification_status": "pending"},
            columns=[_col("phone_number")],
        )
        wb = _read_workbook(xlsx_bytes)
        rows = _data_rows(wb)
        assert len(rows) == 1 + 2  # 2 pending

    def test_q_substring_search_against_phone_number(self, svc, seeded_phones):
        xlsx_bytes, _ = svc.export_phones(
            filters={"q": "2222"},
            columns=[_col("phone_number")],
        )
        wb = _read_workbook(xlsx_bytes)
        rows = _data_rows(wb)
        assert len(rows) == 1 + 1
        assert rows[1][0] == "+15550002222"

    def test_unknown_filter_keys_are_silently_ignored(self, svc, seeded_phones):
        xlsx_bytes, _ = svc.export_phones(
            filters={"future_filter_key": "anything"},
            columns=[_col("phone_number")],
        )
        wb = _read_workbook(xlsx_bytes)
        rows = _data_rows(wb)
        assert len(rows) == 1 + 3


# ===========================================================================
# Empty result + row cap
# ===========================================================================


class TestEmptyAndCap:
    def test_zero_matching_rows_returns_workbook_with_header(self, svc, seeded_phones):
        xlsx_bytes, _ = svc.export_phones(
            filters={"verification_status": "no_such_status"},
            columns=[_col("phone_number", label="Phone")],
        )
        wb = _read_workbook(xlsx_bytes)
        rows = _data_rows(wb)
        assert rows == [("Phone",)]
        assert "export_metadata" in wb.sheetnames

    def test_row_cap_exceeded_raises_value_error(self, svc, session):
        from datetime import datetime, timezone
        from sqlalchemy import insert

        e = Entity(relation_type="primary", deleted_at=not_deleted())
        session.add(e)
        session.flush()
        now = datetime.now(timezone.utc)
        sentinel = not_deleted()
        session.execute(
            insert(PhoneNumber),
            [
                {
                    "entity_id":         e.id,
                    "phone_number":      f"+155500{i:05d}",
                    "ingestion_source":  "manual",
                    "score":             0.0,
                    "deleted_at":        sentinel,
                }
                for i in range(MAX_EXPORT_ROWS + 1)
            ],
        )
        session.commit()

        with pytest.raises(ValueError, match="Too many rows"):
            svc.export_phones(filters={}, columns=[_col("phone_number")])


# ===========================================================================
# Allowlist guard
# ===========================================================================


class TestAllowlistGuard:
    def test_disallowed_key_rejects_whole_request(self, svc, seeded_phones):
        with pytest.raises(ValueError, match="not allowed for export"):
            svc.export_phones(
                filters={},
                columns=[
                    _col("phone_number"),
                    _col("extra_data.api_key"),
                ],
            )

    def test_unknown_flat_key_rejects(self, svc, seeded_phones):
        with pytest.raises(ValueError, match="not allowed for export"):
            svc.export_phones(
                filters={},
                columns=[_col("password")],
            )


# ===========================================================================
# Tasks export
# ===========================================================================


class TestExportTasks:
    def test_happy_path_via_task_service_seed(self, svc, session, seeded_target):
        task_svc = PipelineTaskService(storage=SqlStorage(session))
        task_svc.open_task(phone_id=seeded_target.id, task_type="review")
        task_svc.open_task(phone_id=seeded_target.id, task_type="audit")

        xlsx_bytes, filename = svc.export_tasks(
            filters={},
            columns=[_col("task_type"), _col("status")],
        )
        wb = _read_workbook(xlsx_bytes)
        rows = _data_rows(wb)
        assert len(rows) == 1 + 2
        assert filename.startswith("tasks_")

    def test_exclude_terminal_hides_resolved(self, svc, session, seeded_target):
        task_svc = PipelineTaskService(storage=SqlStorage(session))
        t1 = task_svc.open_task(phone_id=seeded_target.id, task_type="a")
        task_svc.open_task(phone_id=seeded_target.id, task_type="b")
        task_svc.resolve_task(task_id=t1.id, operator_id="adm", outcome="done")

        xlsx_bytes, _ = svc.export_tasks(
            filters={"exclude_terminal": True},
            columns=[_col("task_type")],
        )
        wb = _read_workbook(xlsx_bytes)
        rows = _data_rows(wb)
        assert len(rows) == 1 + 1

    def test_task_allowlist_smoke(self, svc, session, seeded_target):
        task_svc = PipelineTaskService(storage=SqlStorage(session))
        task_svc.open_task(phone_id=seeded_target.id, task_type="review")
        cols = [_col(k) for k in ALLOWED_EXPORT_COLUMNS_TASKS]
        xlsx_bytes, _ = svc.export_tasks(filters={}, columns=cols)
        assert xlsx_bytes[:2] == b"PK"


# ===========================================================================
# Filename hint sanitization
# ===========================================================================


class TestFilenameHint:
    def test_safe_chars_pass_through(self, svc, seeded_phones):
        _, filename = svc.export_phones(
            filters={}, columns=[_col("phone_number")],
            filename_hint="pending_audit",
        )
        assert "pending_audit" in filename

    def test_unsafe_chars_replaced_with_underscores(self, svc, seeded_phones):
        _, filename = svc.export_phones(
            filters={}, columns=[_col("phone_number")],
            filename_hint="test/data export",
        )
        assert "test" in filename
        assert "/" not in filename
        assert " " not in filename


# ===========================================================================
# Metadata sheet
# ===========================================================================


class TestMetadataSheet:
    def test_metadata_sheet_records_filters_and_uuid(self, svc, seeded_phones):
        xlsx_bytes, _ = svc.export_phones(
            filters={"verification_status": "pending"},
            columns=[_col("phone_number")],
        )
        wb = _read_workbook(xlsx_bytes)
        meta = wb["export_metadata"]
        rows_flat = [tuple(c.value for c in row) for row in meta.iter_rows()]
        keys = {r[0] for r in rows_flat}
        assert "bulk_export_id" in keys
        assert "total_rows" in keys
        assert "exported_at_utc" in keys
        assert any("verification_status" in str(r[0]) for r in rows_flat)
