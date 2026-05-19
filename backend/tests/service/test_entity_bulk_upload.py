"""
Service-layer tests for `EntityIngestionService.ingest_bulk_upload`
and `EntityIngestionService.generate_template_xlsx`.

Covers:
    - Excel happy path: parsing + ingestion + audit-trail stamping
    - CSV happy path (UTF-8 BOM tolerance)
    - per-row failures (missing first_name, bad relation, bad target FK)
    - file-shape errors (wrong extension, missing required columns,
      empty workbook) raise ValueError → 422 at the endpoint
    - row-cap enforcement
    - template generation: three sheets, header columns, live targets
"""

import io

import openpyxl
import pytest
from openpyxl import Workbook

from models.entity import Entity
from services.entity_ingestion import (
    BULK_ENTITY_ALL_COLUMNS,
    BULK_ENTITY_REQUIRED_COLUMNS,
    EntityIngestionService,
)


# ===========================================================================
# Fixtures
# ===========================================================================


@pytest.fixture()
def svc(session):
    return EntityIngestionService(session=session)


@pytest.fixture()
def root_target(session):
    e = Entity(entity_type="target", relation_type="primary", client_id=1)
    session.add(e)
    session.commit()
    session.refresh(e)
    return e


@pytest.fixture()
def second_root_target(session):
    e = Entity(entity_type="target", relation_type="primary", client_id=2)
    session.add(e)
    session.commit()
    session.refresh(e)
    return e


def _build_xlsx(header: list, rows: list[list]) -> bytes:
    """Build a one-sheet xlsx with the given header + rows, return bytes."""
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ===========================================================================
# Excel happy path
# ===========================================================================


class TestExcelHappyPath:
    def test_two_rows_ingest_cleanly(
        self, svc, root_target, second_root_target, session
    ):
        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[
                ["Jane", "family",    root_target.id,        "Doe"],
                ["Sam",  "colleague", second_root_target.id, "Chen"],
            ],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        assert summary["success_count"] == 2
        assert summary["failed_count"] == 0
        assert len(summary["entity_ids"]) == 2

        # Each row inherits its OWN target's client_id (independent
        # context per row — distinct from bulk-text's shared default).
        ents = [session.get(Entity, eid) for eid in summary["entity_ids"]]
        clients = sorted(e.client_id for e in ents)
        assert clients == [1, 2]

    def test_audit_trail_stamped_on_every_entity(
        self, svc, root_target, session
    ):
        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[["Jane", "family", root_target.id, "Doe"]],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        ent = session.get(Entity, summary["entity_ids"][0])
        assert ent.extra_data["bulk_submission_id"] == summary["bulk_submission_id"]

    def test_optional_last_name_omitted(self, svc, root_target, session):
        # last_name is optional — operators who paste single-name rows
        # should not see a key inserted into extra_data.
        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[["Cher", "family", root_target.id, None]],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        ent = session.get(Entity, summary["entity_ids"][0])
        assert "last_name" not in ent.extra_data


# ===========================================================================
# CSV happy path
# ===========================================================================


class TestCsvHappyPath:
    def test_utf8_bom_tolerated(self, svc, root_target, session):
        csv_text = (
            "first_name,last_name,relation_type,target_entity_id\n"
            f"Jane,Doe,family,{root_target.id}\n"
        )
        payload = ("﻿" + csv_text).encode("utf-8")    # explicit BOM
        summary = svc.ingest_bulk_upload(payload, "upload.csv")
        assert summary["success_count"] == 1
        ent = session.get(Entity, summary["entity_ids"][0])
        assert ent.extra_data["first_name"] == "Jane"


# ===========================================================================
# Per-row failures
# ===========================================================================


class TestPerRowFailures:
    def test_missing_first_name(self, svc, root_target):
        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[
                ["Jane", "family", root_target.id, "Doe"],
                ["",     "family", root_target.id, "Doe"],   # bad
            ],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        assert summary["success_count"] == 1
        assert summary["failed_count"] == 1
        assert summary["failed_rows"][0]["row"] == 2

    def test_invalid_relation_type(self, svc, root_target):
        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[["Jane", "hacker", root_target.id, "Doe"]],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        assert summary["success_count"] == 0
        assert summary["failed_count"] == 1
        assert "hacker" in summary["failed_rows"][0]["error"]

    def test_disallowed_target_relation(self, svc, root_target):
        # 'target' is in the full vocab but NOT in the operator-creatable
        # subset; bulk-upload must reject it per-row.
        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[["Jane", "target", root_target.id, "Doe"]],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        assert summary["success_count"] == 0
        assert summary["failed_count"] == 1

    def test_disallowed_envelope_relation(self, svc, root_target):
        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[["Jane", "social_envelope", root_target.id, "Doe"]],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        assert summary["success_count"] == 0
        assert summary["failed_count"] == 1

    def test_missing_target(self, svc, root_target):
        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[
                ["Jane", "family", root_target.id, "Doe"],
                ["Sam",  "family", None,           "Chen"],     # bad
            ],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        assert summary["success_count"] == 1
        assert summary["failed_count"] == 1
        assert "target_entity_id" in summary["failed_rows"][0]["error"].lower()

    def test_target_not_found(self, svc, root_target):
        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[["Jane", "family", 99_999, "Doe"]],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        assert summary["success_count"] == 0
        assert "not found" in summary["failed_rows"][0]["error"]

    def test_target_not_root(self, svc, root_target, session):
        # Create an associated entity off the root, then point a row at it.
        assoc = Entity(
            entity_type="family",
            relation_type="associated",
            client_id=root_target.client_id,
            target_entity_id=root_target.id,
        )
        session.add(assoc)
        session.commit()
        session.refresh(assoc)

        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[["Jane", "family", assoc.id, "Doe"]],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        assert summary["success_count"] == 0
        assert "not a root" in summary["failed_rows"][0]["error"]

    def test_target_not_integer(self, svc, root_target):
        # Pasting a string into target_entity_id should land as a clear
        # per-row failure, not a generic 'not found'.
        payload = _build_xlsx(
            header=list(BULK_ENTITY_ALL_COLUMNS),
            rows=[["Jane", "family", "abc", "Doe"]],
        )
        summary = svc.ingest_bulk_upload(payload, "upload.xlsx")
        assert summary["success_count"] == 0
        assert "integer" in summary["failed_rows"][0]["error"].lower()


# ===========================================================================
# File-shape errors (ValueError → endpoint maps to 422)
# ===========================================================================


class TestFileShapeErrors:
    def test_wrong_extension(self, svc):
        with pytest.raises(ValueError, match="Unsupported file extension"):
            svc.ingest_bulk_upload(b"anything", "upload.txt")

    def test_xlsx_missing_required_columns(self, svc):
        # Header is missing 'target_entity_id'.
        payload = _build_xlsx(
            header=["first_name", "last_name", "relation_type"],
            rows=[["Jane", "Doe", "family"]],
        )
        with pytest.raises(ValueError, match="missing required"):
            svc.ingest_bulk_upload(payload, "upload.xlsx")

    def test_csv_missing_required_columns(self, svc):
        payload = b"first_name,last_name,relation_type\nJane,Doe,family\n"
        with pytest.raises(ValueError, match="missing required"):
            svc.ingest_bulk_upload(payload, "upload.csv")

    def test_empty_xlsx(self, svc):
        # Workbook with no rows at all.
        wb = Workbook()
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(ValueError):
            svc.ingest_bulk_upload(buf.getvalue(), "upload.xlsx")

    def test_invalid_xlsx_bytes(self, svc):
        with pytest.raises(ValueError, match="Not a valid"):
            svc.ingest_bulk_upload(b"not a workbook", "upload.xlsx")

    def test_row_cap_exceeded(self, svc, root_target):
        # Build a sheet just over the cap to trip the guard fast.
        from services.entity_ingestion import BULK_ENTITY_MAX_ROWS
        rows = [
            ["Jane", "family", root_target.id, "Doe"]
            for _ in range(BULK_ENTITY_MAX_ROWS + 1)
        ]
        payload = _build_xlsx(list(BULK_ENTITY_ALL_COLUMNS), rows)
        with pytest.raises(ValueError, match="Too many rows"):
            svc.ingest_bulk_upload(payload, "upload.xlsx")


# ===========================================================================
# Template generation
# ===========================================================================


class TestTemplateGeneration:
    def test_template_has_expected_sheets(self, svc):
        payload = svc.generate_template_xlsx()
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True)
        names = wb.sheetnames
        assert "data" in names
        assert "valid_targets" in names
        assert "instructions" in names

    def test_data_sheet_header_matches_required_columns(self, svc):
        payload = svc.generate_template_xlsx()
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True)
        ws = wb["data"]
        header = [c.value for c in next(ws.iter_rows(max_row=1))]
        # Required columns must all appear in the header; ordering is
        # not contractual but optional columns may follow required ones.
        for col in BULK_ENTITY_REQUIRED_COLUMNS:
            assert col in header

    def test_valid_targets_sheet_lists_current_root_targets(
        self, svc, root_target, second_root_target
    ):
        payload = svc.generate_template_xlsx()
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True)
        ws = wb["valid_targets"]
        rows = [tuple(c.value for c in row) for row in ws.iter_rows()]
        # Header row + one row per target.
        assert rows[0] == ("target_entity_id", "client_id")
        ids = {r[0] for r in rows[1:]}
        assert root_target.id in ids
        assert second_root_target.id in ids

    def test_valid_targets_sheet_excludes_associated_entities(
        self, svc, root_target, session
    ):
        # An associated entity should NOT appear in the reference sheet
        # (only root targets are valid `target_entity_id` values).
        assoc = Entity(
            entity_type="family",
            relation_type="associated",
            client_id=root_target.client_id,
            target_entity_id=root_target.id,
        )
        session.add(assoc)
        session.commit()
        session.refresh(assoc)

        payload = svc.generate_template_xlsx()
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True)
        ws = wb["valid_targets"]
        ids = {row[0].value for row in ws.iter_rows(min_row=2)}
        assert root_target.id in ids
        assert assoc.id not in ids

    def test_template_round_trips_through_upload(
        self, svc, root_target, session
    ):
        """
        The generated template should be parseable by the upload path
        once the operator replaces the example target id with a real one.
        We simulate that by editing the example rows in-place.
        """
        payload = svc.generate_template_xlsx()
        # Open, swap the placeholder target_entity_id (0) for a real one,
        # save back to bytes, re-upload.
        wb = openpyxl.load_workbook(io.BytesIO(payload))
        ws = wb["data"]
        header = [c.value for c in next(ws.iter_rows(max_row=1))]
        tgt_col_idx = header.index("target_entity_id") + 1  # 1-based
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=tgt_col_idx, value=root_target.id)
        buf = io.BytesIO()
        wb.save(buf)
        edited = buf.getvalue()

        summary = svc.ingest_bulk_upload(edited, "upload.xlsx")
        assert summary["success_count"] >= 1
        assert summary["failed_count"] == 0
