"""
Unit tests for the file parsers + template generator inside
`services/bulk_ingestion.py`. These tests don't need a database — they
pin the parsing contract that the service-layer + API tests depend on.

Covered:
    - CSV parsing (UTF-8 + BOM-prefixed, header validation, extra cols)
    - XLSX parsing (header-row detection, value coercion, blank rows)
    - Template generation round-trip (write -> re-read -> required cols
      present, instructions sheet exists, sample rows parseable)
"""

import io

import openpyxl
import pytest
from openpyxl import Workbook

from services.bulk_ingestion import (
    BULK_UPLOAD_ALL_COLUMNS,
    BULK_UPLOAD_REQUIRED_COLUMNS,
    BulkIngestionService,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _build_csv(rows: list[list[str]]) -> bytes:
    """Render `rows` (list-of-lists, header first) as UTF-8 CSV bytes."""
    import csv
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def _build_xlsx(rows: list[list]) -> bytes:
    """Render `rows` as an in-memory .xlsx workbook (single sheet)."""
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Required columns are: phone_number, entity_id, ingestion_source
_REQ = list(BULK_UPLOAD_REQUIRED_COLUMNS)  # ["phone_number", "entity_id", "ingestion_source"]


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


class TestParseCsv:
    def test_parses_required_columns(self):
        data = _build_csv([
            _REQ,
            ["+14155551111", "ent-001", "manual"],
            ["+14155551112", "ent-002", "automated"],
        ])
        out = BulkIngestionService._parse_csv(data)
        assert len(out) == 2
        assert out[0]["phone_number"] == "+14155551111"
        assert out[0]["entity_id"] == "ent-001"
        assert out[1]["ingestion_source"] == "automated"

    def test_includes_optional_columns_when_present(self):
        data = _build_csv([
            list(BULK_UPLOAD_ALL_COLUMNS),
            ["+14155551111", "ent-001", "manual", "mobile", "ent-root"],
        ])
        out = BulkIngestionService._parse_csv(data)
        assert out[0]["phone_type"] == "mobile"
        assert out[0]["target_entity_id"] == "ent-root"

    def test_bom_prefixed_utf8_decodes_cleanly(self):
        # Excel-exported CSVs often start with a UTF-8 BOM. The parser
        # uses utf-8-sig so the BOM doesn't leak into the first column
        # name (which would silently fail the required-columns check).
        raw = _build_csv([
            _REQ,
            ["+14155551111", "ent-001", "manual"],
        ])
        with_bom = b"\xef\xbb\xbf" + raw
        out = BulkIngestionService._parse_csv(with_bom)
        assert out[0]["phone_number"] == "+14155551111"

    def test_missing_required_column_raises(self):
        # "ingestion_source" omitted from header.
        data = _build_csv([
            ["phone_number", "entity_id"],
            ["+14155551111", "ent-001"],
        ])
        with pytest.raises(ValueError, match="missing required columns"):
            BulkIngestionService._parse_csv(data)

    def test_empty_csv_raises(self):
        with pytest.raises(ValueError):
            BulkIngestionService._parse_csv(b"")

    def test_extra_columns_preserved_not_rejected(self):
        # An operator's sheet may carry trailing notes columns; the
        # parser is permissive — it returns whatever the header declares.
        data = _build_csv([
            _REQ + ["operator_notes"],
            ["+14155551111", "ent-001", "manual", "VIP customer"],
        ])
        out = BulkIngestionService._parse_csv(data)
        assert out[0].get("operator_notes") == "VIP customer"


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


class TestParseXlsx:
    def test_parses_required_columns(self):
        data = _build_xlsx([
            _REQ,
            ["+14155551111", "ent-001", "manual"],
            ["+14155551112", "ent-002", "automated"],
        ])
        out = BulkIngestionService._parse_xlsx(data)
        assert len(out) == 2
        assert out[0]["phone_number"] == "+14155551111"
        assert out[0]["entity_id"] == "ent-001"

    def test_tolerates_leading_blank_rows(self):
        # Some operators paste a blank first row by accident; the parser
        # treats the first non-empty row as the header.
        data = _build_xlsx([
            [None, None, None],
            _REQ,
            ["+14155551111", "ent-001", "manual"],
        ])
        out = BulkIngestionService._parse_xlsx(data)
        assert len(out) == 1

    def test_skips_blank_data_rows(self):
        data = _build_xlsx([
            _REQ,
            ["+14155551111", "ent-001", "manual"],
            [None, None, None],
            ["+14155551112", "ent-001", "manual"],
        ])
        out = BulkIngestionService._parse_xlsx(data)
        assert len(out) == 2

    def test_missing_required_column_raises(self):
        data = _build_xlsx([
            ["phone_number", "entity_id"],   # no ingestion_source
            ["+14155551111", "ent-001"],
        ])
        with pytest.raises(ValueError, match="missing required columns"):
            BulkIngestionService._parse_xlsx(data)

    def test_invalid_xlsx_bytes_raises(self):
        with pytest.raises(ValueError, match="Not a valid"):
            BulkIngestionService._parse_xlsx(b"not a workbook")

    def test_empty_workbook_raises(self):
        wb = Workbook()
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(ValueError, match="empty"):
            BulkIngestionService._parse_xlsx(buf.getvalue())

    def test_string_values_get_stripped(self):
        # Operators occasionally paste cells with trailing whitespace.
        data = _build_xlsx([
            _REQ,
            ["  +14155551111  ", "  ent-001  ", " manual "],
        ])
        out = BulkIngestionService._parse_xlsx(data)
        assert out[0]["phone_number"] == "+14155551111"
        assert out[0]["entity_id"] == "ent-001"


# ---------------------------------------------------------------------------
# Template generation
# ---------------------------------------------------------------------------


class TestGenerateTemplate:
    def test_returns_nonempty_xlsx_bytes(self):
        payload = BulkIngestionService.generate_template_xlsx()
        assert isinstance(payload, bytes)
        assert len(payload) > 0
        # PK zip signature (xlsx is a zip): first 2 bytes are "PK".
        assert payload[:2] == b"PK"

    def test_template_data_sheet_header_matches_contract(self):
        payload = BulkIngestionService.generate_template_xlsx()
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        ws = wb["data"]
        header_row = next(ws.iter_rows(values_only=True))
        assert list(header_row) == list(BULK_UPLOAD_ALL_COLUMNS)

    def test_template_has_instructions_sheet(self):
        payload = BulkIngestionService.generate_template_xlsx()
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        assert "instructions" in wb.sheetnames

    def test_template_example_rows_round_trip_through_parser(self):
        # The example rows the template ships with should be valid input
        # for `_parse_xlsx` — otherwise we'd be shipping a broken example.
        payload = BulkIngestionService.generate_template_xlsx()
        out = BulkIngestionService._parse_xlsx(payload)
        # Example rows live on the data sheet (under the header).
        assert len(out) >= 1
        for r in out:
            assert r["phone_number"]
            assert r["entity_id"]
            assert r["ingestion_source"]
